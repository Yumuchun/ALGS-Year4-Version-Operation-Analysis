from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
PORTFOLIO = ROOT / "portfolio"
OUT = PORTFOLIO / "figures"
ANALYSIS = PORTFOLIO / "analysis_outputs"
MARKDOWN_PATH = PORTFOLIO / "ALGS_Year4_portfolio_report.md"
DOCX_PATH = PORTFOLIO / "ALGS_Year4_portfolio_report.docx"

EVENT_ORDER = [
    "playoffs_split1",
    "midseason_ewc_split1",
    "playoffs_split2",
    "championship_split2",
]
EVENT_LABEL = {
    "playoffs_split1": "Split 1",
    "midseason_ewc_split1": "EWC",
    "playoffs_split2": "Split 2",
    "championship_split2": "Champs",
}

FONT_REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def canvas(title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1400, 820), "white")
    draw = ImageDraw.Draw(img)
    draw.text((72, 44), title, fill="#111827", font=font(34, True))
    if subtitle:
        draw.text((72, 92), subtitle, fill="#4B5563", font=font(20))
    return img, draw


def draw_axes(draw: ImageDraw.ImageDraw, x0: int, y0: int, x1: int, y1: int, max_y: float, ticks: int = 5):
    draw.line((x0, y1, x1, y1), fill="#9CA3AF", width=2)
    draw.line((x0, y0, x0, y1), fill="#9CA3AF", width=2)
    for i in range(ticks + 1):
        y = y1 - (y1 - y0) * i / ticks
        val = max_y * i / ticks
        draw.line((x0 - 6, y, x1, y), fill="#EEF2F7", width=1)
        draw.text((x0 - 56, y - 10), f"{val:.0f}", fill="#6B7280", font=font(16), anchor=None)


def y_map(value: float, y0: int, y1: int, max_y: float) -> float:
    return y1 - (value / max_y) * (y1 - y0)


def draw_legend(draw: ImageDraw.ImageDraw, items: list[tuple[str, str]], x: int, y: int):
    cursor = x
    for label, color in items:
        draw.rounded_rectangle((cursor, y, cursor + 28, y + 16), radius=4, fill=color)
        draw.text((cursor + 38, y - 3), label, fill="#374151", font=font(17))
        cursor += 38 + int(draw.textlength(label, font=font(17))) + 32


def chart_score_structure():
    df = pd.read_csv(ANALYSIS / "score_response_by_event.csv")
    df = df.set_index("event_slug").loc[EVENT_ORDER].reset_index()
    reset = pd.read_csv(ANALYSIS / "reset_signal_by_event.csv")
    reset = reset.set_index("event_slug").loc[EVENT_ORDER].reset_index()
    labels = [EVENT_LABEL[e] for e in df["event_slug"]]
    kill_share = (df["kill_point_share"] * 100).tolist()
    score_kills = reset["score_kills_per_game"].tolist()
    deaths = reset["deaths_per_game"].tolist()
    respawns = reset["respawns_per_game"].tolist()
    revives = reset["revives_per_game"].tolist()

    img, draw = canvas(
        "Score Stayed Stable, Reset Signals Rose",
        "Credited kills stayed near the BR baseline; later metas created more respawn and revive windows.",
    )
    left = (72, 164, 674, 610)
    right = (742, 164, 1328, 610)

    def panel(box: tuple[int, int, int, int], title: str) -> None:
        x0, y0, x1, y1 = box
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill="#F9FAFB", outline="#E5E7EB", width=2)
        draw.text((x0 + 24, y0 + 24), title, fill="#111827", font=font(22, True))

    def y_range(value: float, y0: int, y1: int, min_y: float, max_y: float) -> float:
        return y1 - ((value - min_y) / (max_y - min_y)) * (y1 - y0)

    def line_chart(
        box: tuple[int, int, int, int],
        series: list[tuple[str, list[float], str]],
        min_y: float,
        max_y: float,
        ticks: list[float],
    ) -> None:
        x0, y0, x1, y1 = box
        plot_x0, plot_y0, plot_x1, plot_y1 = x0 + 76, y0 + 94, x1 - 36, y1 - 66
        for tick in ticks:
            y = y_range(tick, plot_y0, plot_y1, min_y, max_y)
            draw.line((plot_x0, y, plot_x1, y), fill="#E5E7EB", width=1)
            draw.text((x0 + 26, y - 10), f"{tick:.0f}", fill="#9CA3AF", font=font(16))
        xs = [plot_x0 + (plot_x1 - plot_x0) * i / (len(labels) - 1) for i in range(len(labels))]
        for legend_idx, (legend, values, color) in enumerate(series):
            points = [(xs[i], y_range(values[i], plot_y0, plot_y1, min_y, max_y), values[i]) for i in range(len(labels))]
            for a, b in zip(points, points[1:]):
                draw.line((a[0], a[1], b[0], b[1]), fill=color, width=5)
            for x, y, value in points:
                draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=color, outline="white", width=3)
                draw.text((x, y - 24), f"{value:.1f}", fill=color, font=font(16, True), anchor="mm")
            lx = plot_x0 + legend_idx * 142
            draw.rounded_rectangle((lx, y0 + 60, lx + 22, y0 + 74), radius=4, fill=color)
            draw.text((lx + 30, y0 + 54), legend, fill="#4B5563", font=font(16))
        for x, label in zip(xs, labels):
            draw.text((x, plot_y1 + 28), label, fill="#111827", font=font(18, True), anchor="mm")

    def grouped_bars(box: tuple[int, int, int, int]) -> None:
        x0, y0, x1, y1 = box
        plot_x0, plot_y0, plot_x1, plot_y1 = x0 + 76, y0 + 94, x1 - 36, y1 - 66
        for tick in [0, 5, 10, 15, 20]:
            y = y_range(tick, plot_y0, plot_y1, 0, 22)
            draw.line((plot_x0, y, plot_x1, y), fill="#E5E7EB", width=1)
            draw.text((x0 + 26, y - 10), f"{tick:.0f}", fill="#9CA3AF", font=font(16))
        group_gap = (plot_x1 - plot_x0) / len(labels)
        bar_w = 34
        for i, label in enumerate(labels):
            cx = plot_x0 + group_gap * (i + 0.5)
            for offset, value, color in [(-22, respawns[i], "#F59E0B"), (22, revives[i], "#10B981")]:
                x = cx + offset - bar_w / 2
                y = y_range(value, plot_y0, plot_y1, 0, 22)
                draw.rounded_rectangle((x, y, x + bar_w, plot_y1), radius=6, fill=color)
                draw.text((x + bar_w / 2, y - 16), f"{value:.1f}", fill=color, font=font(16, True), anchor="mm")
            draw.text((cx, plot_y1 + 28), label, fill="#111827", font=font(18, True), anchor="mm")
        draw.rounded_rectangle((plot_x0, y0 + 60, plot_x0 + 22, y0 + 74), radius=4, fill="#F59E0B")
        draw.text((plot_x0 + 30, y0 + 54), "Respawn (rspn)", fill="#4B5563", font=font(16))
        draw.rounded_rectangle((plot_x0 + 176, y0 + 60, plot_x0 + 198, y0 + 74), radius=4, fill="#10B981")
        draw.text((plot_x0 + 206, y0 + 54), "Revive (rez)", fill="#4B5563", font=font(16))

    panel(left, "A. Credited kills vs deaths")
    panel(right, "B. Reset indicators per game")
    line_chart(left, [("Score kills", score_kills, "#2563EB"), ("Deaths", deaths, "#EF4444")], 54, 62, [54, 56, 58, 60, 62])
    grouped_bars(right)

    draw.text((72, 650), "Readout:", fill="#111827", font=font(21, True))
    draw.text(
        (172, 650),
        "Kill share stayed around 51%-52%, but Split 2 / Championship show more second-chance space.",
        fill="#374151",
        font=font(19),
    )
    draw.text(
        (72, 686),
        "Data: match_scores.csv for credited kills and score share; game_team_stats.csv for deaths, rspn and rez.",
        fill="#6B7280",
        font=font(17),
    )
    draw.text(
        (72, 716),
        "Key caution: deaths, respawns and revives are reset signals, not extra scoring kills.",
        fill="#6B7280",
        font=font(17),
    )
    img.save(OUT / "score_structure.png", quality=95)


def chart_composition_concentration():
    df = pd.read_csv(ANALYSIS / "composition_concentration_by_event.csv")
    df = df.set_index("event_slug").loc[EVENT_ORDER].reset_index()
    labels = [EVENT_LABEL[e] for e in df["event_slug"]]
    series = [
        ("Top 1 comp", "top1_pick_rate_pct", "#2563EB"),
        ("Top 3 comps", "top3_pick_rate_pct_sum", "#10B981"),
        ("Top 5 comps", "top5_pick_rate_pct_sum", "#F59E0B"),
    ]
    img, draw = canvas(
        "Composition Concentration Rebuilt After Every Patch",
        "The meta was loosened at EWC, then re-converged sharply by Championship.",
    )
    x0, y0, x1, y1 = 120, 170, 1280, 680
    draw_axes(draw, x0, y0, x1, y1, 100)
    group_w = 210
    bar_w = 52
    gap = (x1 - x0) / len(labels)
    for i, label in enumerate(labels):
        cx = x0 + gap * (i + 0.5)
        for j, (_, col, color) in enumerate(series):
            x = cx - group_w / 2 + j * (bar_w + 18)
            val = float(df.loc[i, col])
            y = y_map(val, y0, y1, 100)
            draw.rounded_rectangle((x, y, x + bar_w, y1), radius=7, fill=color)
            draw.text((x + bar_w / 2, y - 15), f"{val:.0f}", fill="#374151", font=font(14), anchor="mm")
        draw.text((cx, y1 + 24), label, fill="#111827", font=font(19, True), anchor="mm")
    draw_legend(draw, [(name, color) for name, _, color in series], 120, 715)
    img.save(OUT / "composition_concentration.png", quality=95)


def chart_weapon_shift():
    df = pd.read_csv(ANALYSIS / "weapon_category_by_event.csv")
    keep = ["AR", "SMG", "Precision", "Shotgun", "Shotgun/Akimbo", "LMG"]
    colors = {
        "AR": "#2563EB",
        "SMG": "#7C3AED",
        "Precision": "#0891B2",
        "Shotgun": "#DC2626",
        "Shotgun/Akimbo": "#F97316",
        "LMG": "#16A34A",
    }
    pivot = (
        df[df["category"].isin(keep)]
        .pivot_table(index="event_slug", columns="category", values="kill_share_pct", aggfunc="sum")
        .reindex(EVENT_ORDER)
        .fillna(0)
    )
    labels = [EVENT_LABEL[e] for e in pivot.index]
    img, draw = canvas(
        "Weapon Kill Share Shifted Between Weapon Ecosystems",
        "AR dominance gave way to Akimbo/Shotgun, then to Shotgun + Precision at Championship.",
    )
    x0, y0, x1, y1 = 120, 170, 1280, 680
    draw_axes(draw, x0, y0, x1, y1, 100)
    gap = (x1 - x0) / len(labels)
    bar_w = 150
    for i, event in enumerate(pivot.index):
        cx = x0 + gap * (i + 0.5)
        bottom = y1
        for cat in keep:
            val = float(pivot.loc[event, cat])
            h = (val / 100) * (y1 - y0)
            if h > 0:
                draw.rectangle((cx - bar_w / 2, bottom - h, cx + bar_w / 2, bottom), fill=colors[cat])
                if h >= 36:
                    draw.text((cx, bottom - h / 2), f"{cat}\n{val:.0f}%", fill="white", font=font(15, True), anchor="mm", align="center")
            bottom -= h
        draw.rectangle((cx - bar_w / 2, y0, cx + bar_w / 2, y1), outline="#E5E7EB", width=1)
        draw.text((cx, y1 + 24), labels[i], fill="#111827", font=font(19, True), anchor="mm")
    draw_legend(draw, [(cat, colors[cat]) for cat in keep], 120, 715)
    img.save(OUT / "weapon_shift.png", quality=95)


def chart_class_slots():
    df = pd.read_csv(ANALYSIS / "legend_class_share_by_event.csv")
    classes = ["Assault", "Controller", "Recon", "Support", "Skirmisher", "Other"]
    colors = {
        "Assault": "#2563EB",
        "Controller": "#10B981",
        "Recon": "#8B5CF6",
        "Support": "#F97316",
        "Skirmisher": "#06B6D4",
        "Other": "#9CA3AF",
    }
    pivot = (
        df.pivot_table(index="event_slug", columns="class", values="pick_rate_pct_sum", aggfunc="sum")
        .reindex(EVENT_ORDER)
        .reindex(columns=classes)
        .fillna(0)
    )
    labels = [EVENT_LABEL[e] for e in pivot.index]
    img, draw = canvas(
        "Legend Class Slots Show What the Meta Needed",
        "Class share sums three legend slots per team; 200% Support means roughly two Support slots per team.",
    )
    x0, y0, x1, y1 = 120, 170, 1280, 680
    draw_axes(draw, x0, y0, x1, y1, 300)
    gap = (x1 - x0) / len(labels)
    bar_w = 150
    for i, event in enumerate(pivot.index):
        cx = x0 + gap * (i + 0.5)
        bottom = y1
        for cls in classes:
            val = float(pivot.loc[event, cls])
            h = (val / 300) * (y1 - y0)
            if h > 0:
                draw.rectangle((cx - bar_w / 2, bottom - h, cx + bar_w / 2, bottom), fill=colors[cls])
                if h >= 40:
                    draw.text((cx, bottom - h / 2), f"{cls}\n{val:.0f}%", fill="white", font=font(15, True), anchor="mm", align="center")
            bottom -= h
        draw.rectangle((cx - bar_w / 2, y0, cx + bar_w / 2, y1), outline="#E5E7EB", width=1)
        draw.text((cx, y1 + 24), labels[i], fill="#111827", font=font(19, True), anchor="mm")
    draw_legend(draw, [(cls, colors[cls]) for cls in classes], 120, 715)
    img.save(OUT / "legend_class_slots.png", quality=95)


def chart_style_ppg():
    df = pd.read_csv(ANALYSIS / "team_style_summary.csv")
    styles = [
        ("edge_fighting_proxy", "Edge"),
        ("hybrid_high_yield", "Hybrid"),
        ("zone_control_proxy", "Zone"),
        ("low_yield_or_unstable", "Low/unstable"),
    ]
    colors = {
        "edge_fighting_proxy": "#DC2626",
        "hybrid_high_yield": "#2563EB",
        "zone_control_proxy": "#10B981",
        "low_yield_or_unstable": "#9CA3AF",
    }
    pivot = (
        df.pivot_table(index="event_slug", columns="style", values="avg_total_ppg", aggfunc="mean")
        .reindex(EVENT_ORDER)
    )
    labels = [EVENT_LABEL[e] for e in pivot.index]
    img, draw = canvas(
        "Pure Edge Fighting Scored Less Than Hybrid or Zone Styles",
        "High-kill games mattered, but the best yield came from turning fights into Top 5 conversions.",
    )
    x0, y0, x1, y1 = 120, 170, 1280, 680
    draw_axes(draw, x0, y0, x1, y1, 8)
    xs = [x0 + (x1 - x0) * (i + 0.5) / len(labels) for i in range(len(labels))]
    for style, label in styles:
        pts = []
        for i, event in enumerate(pivot.index):
            val = float(pivot.loc[event, style])
            pts.append((xs[i], y_map(val, y0, y1, 8), val))
        for a, b in zip(pts, pts[1:]):
            draw.line((a[0], a[1], b[0], b[1]), fill=colors[style], width=5)
        for x, y, val in pts:
            draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill=colors[style], outline="white", width=3)
            draw.text((x, y - 28), f"{val:.1f}", fill=colors[style], font=font(15, True), anchor="mm")
    for x, label in zip(xs, labels):
        draw.text((x, y1 + 24), label, fill="#111827", font=font(19, True), anchor="mm")
    draw_legend(draw, [(label, colors[style]) for style, label in styles], 120, 715)
    draw.text((78, 152), "PPG", fill="#6B7280", font=font(16))
    img.save(OUT / "team_style_ppg.png", quality=95)


def add_runs(paragraph, text: str):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_doc_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Title", 26, "000000", 0, 3),
        ("Heading 1", 20, "000000", 20, 6),
        ("Heading 2", 16, "000000", 18, 6),
        ("Heading 3", 14, "434343", 16, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
            if row_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def markdown_to_docx():
    doc = Document()
    set_doc_styles(doc)
    lines = MARKDOWN_PATH.read_text(encoding="utf-8").splitlines()
    i = 0
    title_done = False

    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith("!["):
            alt = line[line.find("[") + 1 : line.find("]")]
            path_text = line[line.find("(") + 1 : line.rfind(")")]
            image_path = (PORTFOLIO / path_text).resolve() if not Path(path_text).is_absolute() else Path(path_text)
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path), width=Inches(6.3))
                if alt:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(10)
                    run = cap.add_run(alt)
                    run.italic = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(85, 85, 85)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for idx, table_line in enumerate(table_lines):
                if idx == 1 and re.match(r"^\|[\s:\-|]+\|$", table_line):
                    continue
                cells = [c.strip() for c in table_line.strip("|").split("|")]
                rows.append(cells)
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], cell_text)
                style_table(table)
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title" if not title_done else "Heading 1")
            add_runs(p, line[2:].strip())
            title_done = True
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_runs(p, line[3:].strip())
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_runs(p, line[5:].strip())
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\. ", "", line).strip())
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    doc.save(DOCX_PATH)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    chart_score_structure()
    chart_composition_concentration()
    chart_weapon_shift()
    chart_class_slots()
    chart_style_ppg()
    if MARKDOWN_PATH.exists():
        markdown_to_docx()


if __name__ == "__main__":
    main()
